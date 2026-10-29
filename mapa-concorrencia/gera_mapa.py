"""Gera o mapa de concorrência (Excel) de uma cotação.

Uso:
    python gera_mapa.py exemplo-entrada.json [saida.xlsx]

Sem o segundo argumento, grava no caminho do campo "saida" do JSON.

Entrada (JSON), um objeto com:
  codigo, obra, endereco   textos do cabeçalho da planilha
  itens                    lista de {id, seq, descricao, quantidade, unidade}: o que foi pedido
  fornecedores             lista de objetos, cada um com
      nome, telefone, contato
      precos               {"<id do item>": {"unit": 0.0, "total": 0.0}}
      extras               lista de textos: itens ofertados fora do pedido
      valor_imposto, obs_imposto, frete, valor_frete
      pct_desconto         em %, ou seja 5 vale 5%
      valor_desconto, observacoes, forma_pagamento
      prazo_servico, prazo_entrega, validade

Números podem vir como texto: com vírgula valem como decimal brasileiro ("5,5", "1.234,50");
texto que não é número dá erro (nunca vira "não cotou" calado).

O gerador também aceita o payload que o fluxo n8n 02 monta (itens com "numero" e
"id_item_pedido"; fornecedor com "itens", "nome_contato" e "prazo_material").

Layout (aba "Cotação"): cabeçalho do fornecedor nas linhas 4 a 6, duas colunas por
fornecedor a partir da E (Unit. / V. Total), um item por linha a partir da 8, e então
Total, Imposto, Frete, Desconto, Total Geral, Itens sem cotação, Observação e o
Resumo de negociação. Quantos itens e quantos fornecedores couberem: as linhas e
as colunas são montadas na hora, sem molde.
"""
import json
import re
import sys
import unicodedata

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as col
from openpyxl.worksheet.properties import PageSetupProperties

PRIMEIRA_COL = 5  # E
PRIMEIRA_LINHA_ITEM = 8
NAO_COTOU = "não cotou"

AZUL = PatternFill("solid", start_color="1F3A5F", end_color="1F3A5F")
CINZA = PatternFill("solid", start_color="E7EAEE", end_color="E7EAEE")
VERDE = PatternFill("solid", start_color="C6EFCE", end_color="C6EFCE")
VERMELHO = PatternFill("solid", start_color="F8CBAD", end_color="F8CBAD")
FINO = Side(style="thin", color="9AA3AD")
BORDA = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)
BRL = '"R$" #,##0.00##'


def num(v):
    """Número vindo do JSON. Texto com vírgula é decimal brasileiro ("5,5", "1.234,50").
    O que não der para ler vira erro, nunca "não cotou" calado: um preço escondido não pode passar batido."""
    if v is None or v == "":
        return None
    bruto = v
    if isinstance(v, str) and "," in v:
        v = v.replace(".", "").replace(",", ".")
    try:
        return float(v)
    except (TypeError, ValueError):
        raise ValueError(f"valor numérico inválido: {bruto!r}") from None


def normaliza(d):
    """Devolve (itens, fornecedores) no formato interno, aceitando os dois contratos."""
    itens = []
    for i, it in enumerate(d.get("itens") or [], start=1):
        itens.append({
            "id": it.get("id", it.get("id_item_pedido")),
            "seq": int(it.get("seq") or it.get("numero") or i),  # "2" e 2 se misturam sem quebrar a ordenação
            "descricao": it.get("descricao") or "",
            "quantidade": num(it.get("quantidade")),
            "unidade": it.get("unidade"),
        })
    itens.sort(key=lambda x: x["seq"])
    forns = []
    for f in d.get("fornecedores") or []:
        g = dict(f)
        g["contato"] = f.get("contato") or f.get("nome_contato")
        g["prazo_entrega"] = f.get("prazo_entrega") or f.get("prazo_material")
        if "precos" not in f:  # contrato do n8n: {"<id>": {"preco_unit": ..}}
            g["precos"] = {
                str(k): {"unit": v.get("preco_unit"), "total": v.get("preco_total")}
                for k, v in (f.get("itens") or {}).items()
            }
        forns.append(g)
    return itens, forns


def imposto_incluso(obs):
    """True se o texto do fornecedor diz que o imposto JÁ está nos preços ("Incluso", "incluído no preço").
    Negação na mesma frase desfaz ("Não incluso", "Impostos não inclusos", "ICMS não está incluso").
    Sem acento e sem caixa, para "NÃO" e "nao" valerem igual.
    Limitação: é uma heurística de palavras; uma frase longa, com vírgula no meio e vários sentidos, pode escapar.
    Na dúvida o imposto SOMA, para o erro nunca esconder custo."""
    s = unicodedata.normalize("NFD", obs or "").encode("ascii", "ignore").decode().lower()
    inclui = r"\binclu(?:s[oa]s?|id[oa]s?|i|indo)\b"
    return bool(re.search(inclui, s)) and not re.search(r"\b(?:nao|sem)\b[^,;.:()/-]*" + inclui, s)


def preco_do_item(precos, it):
    """(unitário, total) que o fornecedor deu para o item. Veio só o total e há quantidade: o unitário sai da conta."""
    informado = precos.get(str(it["id"])) or {}
    unit, total = num(informado.get("unit")), num(informado.get("total"))
    if unit is None and it["quantidade"] and total is not None:
        unit = total / it["quantidade"]
    return unit, total


def cabecalho(c, texto, fill=AZUL):
    c.value = texto
    c.fill = fill
    c.font = Font(bold=True, color="FFFFFF" if fill is AZUL else "000000")
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def mescla(ws, r1, c1, r2, c2):
    ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            ws.cell(r, c).border = BORDA


def gera_mapa(d, destino):
    """Monta o workbook a partir do dicionário `d` e grava em `destino` (caminho ou arquivo)."""
    itens, forns = normaliza(d)
    if not itens:
        raise ValueError("cotação sem itens: não há o que comparar")
    n = len(itens)
    ult = PRIMEIRA_LINHA_ITEM + n - 1
    L_TOTAL, L_IMP, L_FRETE, L_DESC, L_GERAL, L_SEMCOT, L_OBS = (ult + 1 + k for k in range(7))
    L_RESUMO = L_OBS + 2
    L_PAGAMENTO, L_PRAZO_SERVICO, L_PRAZO_MATERIAL, L_VALIDADE, L_OBS_COMPRADOR = (L_RESUMO + 1 + k for k in range(5))
    ult_col = PRIMEIRA_COL + 2 * max(len(forns), 1) - 1

    wb = openpyxl.Workbook()
    wb.properties.creator = "agente-compras-whatsapp"
    ws = wb.active
    ws.title = "Cotação"

    # --- título e dados da obra -------------------------------------------------
    codigo = d.get("codigo") or d.get("codigo_cotacao")
    titulo = "MAPA DE CONCORRÊNCIA" + (f" - {codigo}" if codigo else "")
    for r, texto in enumerate((f"Obra: {d.get('obra') or ''}", f"Endereço: {d.get('endereco') or ''}", titulo), start=1):
        ws.cell(r, 1, texto)
        mescla(ws, r, 1, r, ult_col)
        ws.row_dimensions[r].height = 21
    cabecalho(ws["A3"], titulo)
    ws["A1"].font = ws["A2"].font = Font(bold=True)

    # --- coluna da esquerda: o que foi pedido -------------------------------------
    for c, titulo in enumerate(["Item", "DESCRIÇÃO", "Quantidade", "Unidade"], start=1):
        mescla(ws, 4, c, 7, c)
        cabecalho(ws.cell(4, c), titulo)
    for r, it in enumerate(itens, start=PRIMEIRA_LINHA_ITEM):
        ws.cell(r, 1, it["seq"])
        ws.cell(r, 2, it["descricao"]).alignment = Alignment(wrap_text=True, vertical="center")
        ws.cell(r, 3, it["quantidade"])
        ws.cell(r, 4, it["unidade"])
        for c in range(1, 5):
            ws.cell(r, c).border = BORDA
    rotulos = {L_TOTAL: "Total:", L_IMP: "Imposto", L_FRETE: "Frete", L_DESC: "Desconto",
               L_GERAL: "Total Geral", L_SEMCOT: "Itens sem cotação", L_OBS: "Observação do fornecedor",
               L_PAGAMENTO: "Condições de pagamento", L_PRAZO_SERVICO: "Prazo para execução do serviço",
               L_PRAZO_MATERIAL: "Prazo para entrega do material", L_VALIDADE: "Validade da proposta",
               L_OBS_COMPRADOR: "Observação do comprador"}
    for r, texto in rotulos.items():
        mescla(ws, r, 1, r, 4)
        ws.cell(r, 1, texto).font = Font(bold=True)
    mescla(ws, L_RESUMO, 1, L_RESUMO, ult_col)
    cabecalho(ws.cell(L_RESUMO, 1), "RESUMO DE NEGOCIAÇÃO")

    # cada fornecedor ocupa duas colunas: preço unitário e valor total
    unit_cols = []
    completos = []  # (coluna, total geral) só de quem cotou todos os itens
    for i, f in enumerate(forns):
        cu = PRIMEIRA_COL + 2 * i
        cv = cu + 1
        U, V = col(cu), col(cv)
        unit_cols.append(U)
        ws.column_dimensions[U].width = 15
        ws.column_dimensions[V].width = 17
        for r, valor in ((4, f.get("nome")), (5, f.get("telefone")), (6, f.get("contato"))):
            mescla(ws, r, cu, r, cv)
            ws.cell(r, cu, valor).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.cell(4, cu).font = Font(bold=True)
        ws.row_dimensions[4].height = 36
        cabecalho(ws.cell(7, cu), "Unit.", CINZA)
        cabecalho(ws.cell(7, cv), "V. Total", CINZA)

        precos = f.get("precos") or {}
        avisos, sem_preco, soma = [], 0, 0.0
        for r, it in enumerate(itens, start=PRIMEIRA_LINHA_ITEM):
            unit, total = preco_do_item(precos, it)
            qtd = it["quantidade"]
            if unit is not None:
                ws.cell(r, cu, round(unit, 4))
                if qtd is not None:
                    ws.cell(r, cv, f"={U}{r}*C{r}")
                    soma += unit * qtd
                    if total is not None and abs(total - unit * qtd) > 0.01:
                        avisos.append(f"Item {it['seq']}: total informado R$ {total:.2f} difere de unit. x qtd. (R$ {unit * qtd:.2f})")
                else:
                    valor_linha = unit if total is None else total
                    ws.cell(r, cv, valor_linha)
                    soma += valor_linha
            elif total is not None:
                ws.cell(r, cv, total)
                soma += total
            else:
                ws.cell(r, cu, NAO_COTOU).fill = VERMELHO
                sem_preco += 1
            for c in (cu, cv):
                ws.cell(r, c).border = BORDA
                ws.cell(r, c).number_format = BRL

        # rodapé: total, imposto, frete, desconto, total geral
        for r in (L_TOTAL, L_GERAL, L_SEMCOT, L_OBS, L_PAGAMENTO, L_PRAZO_SERVICO, L_PRAZO_MATERIAL, L_VALIDADE):
            mescla(ws, r, cu, r, cv)
        ws.cell(L_TOTAL, cu, f"=SUM({V}{PRIMEIRA_LINHA_ITEM}:{V}{ult})")

        imp = num(f.get("valor_imposto"))
        obs_imp = f.get("obs_imposto")
        ws.cell(L_IMP, cu, imp or 0)
        ws.cell(L_IMP, cv, obs_imp or ("Não informado" if not imp else None))  # silêncio não vira "Incluso"
        imp_incluso = imposto_incluso(ws.cell(L_IMP, cv).value)  # já está nos preços: não soma de novo

        frete = num(f.get("valor_frete"))
        modalidade_frete = f.get("frete") or (None if frete is not None else "Não informado")
        ws.cell(L_FRETE, cu, frete or 0)
        ws.cell(L_FRETE, cv, modalidade_frete)

        dv, dp = num(f.get("valor_desconto")), num(f.get("pct_desconto"))
        if dv is not None:
            desconto = dv
        elif dp:
            desconto = f"={U}{L_TOTAL}*{round(dp / 100, 6)}"
        else:
            desconto = 0
        ws.cell(L_DESC, cv, desconto)
        ws.cell(L_DESC, cu, f"=IF({U}{L_TOTAL}>0,{V}{L_DESC}/{U}{L_TOTAL},0)").number_format = "0.0%"

        termos = [f"{U}{L_TOTAL}"] + ([] if imp_incluso else [f"{U}{L_IMP}"]) + [f"{U}{L_FRETE}"]
        ws.cell(L_GERAL, cu, "=" + "+".join(termos) + f"-{V}{L_DESC}")
        ws.cell(L_SEMCOT, cu, f"=COUNTIF({U}{PRIMEIRA_LINHA_ITEM}:{U}{ult},\"{NAO_COTOU}\")")
        for r in (L_TOTAL, L_IMP, L_FRETE, L_GERAL):
            ws.cell(r, cu).number_format = BRL
        for r in (L_DESC,):
            ws.cell(r, cv).number_format = BRL
        ws.cell(L_GERAL, cu).font = Font(bold=True)

        # observações: o que o fornecedor escreveu + o que o mapa percebeu
        obs = [f.get("observacoes") or ""]
        if sem_preco:
            obs.append("Não cotou: " + "; ".join(it["descricao"] for it in itens if not cotou(precos, it)))
        extras = f.get("extras") or []
        if extras:
            lista_extras = "; ".join(extras)
            obs.append(f"Itens ofertados fora do pedido: {lista_extras}")
        obs += avisos
        texto_obs = "\n".join(o for o in obs if o) or None
        ws.cell(L_OBS, cu, texto_obs).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(L_PAGAMENTO, cu, f.get("forma_pagamento") or "Não informado no orçamento")
        ws.cell(L_PRAZO_SERVICO, cu, f.get("prazo_servico") or "Não se aplica")
        ws.cell(L_PRAZO_MATERIAL, cu, f.get("prazo_entrega") or "Não informado")
        ws.cell(L_VALIDADE, cu, f.get("validade") or "Não informado")
        for r in (L_PAGAMENTO, L_PRAZO_SERVICO, L_PRAZO_MATERIAL, L_VALIDADE):
            ws.cell(r, cu).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        if sem_preco == 0:
            geral = soma + (0 if imp_incluso else (imp or 0)) + (frete or 0) - (dv if dv is not None else soma * (dp or 0) / 100)
            completos.append((U, geral))

    mescla(ws, L_OBS_COMPRADOR, PRIMEIRA_COL, L_OBS_COMPRADOR, ult_col)  # preenchimento manual

    # --- destaques -------------------------------------------------------------
    # menor preço unitário por item (regra viva: acompanha edição na planilha)
    if len(forns) >= 2:
        for U in unit_cols:
            alvo = ",".join(f"${c}{PRIMEIRA_LINHA_ITEM}" for c in unit_cols)
            ws.conditional_formatting.add(
                f"{U}{PRIMEIRA_LINHA_ITEM}:{U}{ult}",
                FormulaRule(formula=[f"AND(ISNUMBER({U}{PRIMEIRA_LINHA_ITEM}),{U}{PRIMEIRA_LINHA_ITEM}=MIN({alvo}))"], fill=VERDE),
            )
    # menor total geral entre quem cotou TUDO (quem deixou item sem preço não entra na disputa)
    if len(completos) >= 2:
        menor = min(completos, key=lambda x: x[1])[0]
        ws[f"{menor}{L_GERAL}"].fill = VERDE

    # --- finalização -----------------------------------------------------------
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 46
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 10
    ws.row_dimensions[L_OBS].height = 60
    ws.row_dimensions[L_OBS_COMPRADOR].height = 60
    ws.freeze_panes = f"{col(PRIMEIRA_COL)}{PRIMEIRA_LINHA_ITEM}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    wb.save(destino)
    return {"itens": n, "fornecedores": len(forns)}


def cotou(precos, it):
    unit, total = preco_do_item(precos, it)
    return unit is not None or total is not None


def main(argv):
    if len(argv) < 2:
        print(__doc__)
        return 2
    with open(argv[1], encoding="utf-8") as f:
        dados = json.load(f)
    destino = argv[2] if len(argv) > 2 else dados.get("saida")
    if not destino:
        print("informe o arquivo de saída (2º argumento ou campo \"saida\")", file=sys.stderr)
        return 2
    try:
        gera_mapa(dados, destino)
    except ValueError as e:  # dado ruim na entrada: mensagem clara, sem traceback
        print(f"erro: {e}", file=sys.stderr)
        return 1
    print(destino)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
