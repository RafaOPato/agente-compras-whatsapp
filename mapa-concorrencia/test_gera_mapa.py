"""Testes do gerador do mapa. Rodar: python -m unittest -v   (precisa de openpyxl)

O openpyxl não calcula fórmulas, então o teste tem um avaliador mínimo só para as
fórmulas que o gerador escreve (SUM, COUNTIF, IF, + - * /) e compara com uma conta
feita à parte, direto do JSON de entrada.
"""
import http.client
import json
import os
import pathlib
import re
import tempfile
import threading
import unittest
import urllib.error
import urllib.request

import openpyxl
from openpyxl.utils import get_column_letter

import gera_mapa
import servidor_http

PASTA = pathlib.Path(__file__).resolve().parent
ENTRADA = PASTA / "exemplo-entrada.json"


def gera(dados):
    with tempfile.TemporaryDirectory() as tmp:
        caminho = os.path.join(tmp, "mapa.xlsx")
        gera_mapa.gera_mapa(dados, caminho)
        return openpyxl.load_workbook(caminho).active


def avalia(ws, ref):
    """Valor da célula `ref`, resolvendo as fórmulas que o gerador usa."""
    v = ws[ref].value
    if not (isinstance(v, str) and v.startswith("=")):
        return v
    f = v[1:]
    f = re.sub(r"SUM\(([A-Z]+\d+):([A-Z]+\d+)\)", lambda m: str(soma(ws, m.group(1), m.group(2))), f)
    f = re.sub(r'COUNTIF\(([A-Z]+\d+):([A-Z]+\d+),"([^"]+)"\)',
               lambda m: str(sum(1 for c in faixa(ws, m.group(1), m.group(2)) if c.value == m.group(3))), f)
    f = re.sub(r"IF\(([A-Z]+\d+)>0,(.+),0\)$", r"((\2) if \1>0 else 0)", f)
    f = re.sub(r"\b([A-Z]+\d+)\b", lambda m: repr(float(avalia(ws, m.group(1)) or 0)), f)
    return eval(f, {"__builtins__": {}})  # só números e operadores vindos do gerador


def faixa(ws, a, b):
    return [c for linha in ws[a:b] for c in linha]


def soma(ws, a, b):
    return sum(float(avalia(ws, c.coordinate) or 0) for c in faixa(ws, a, b) if not isinstance(c.value, str) or c.value.startswith("="))


def esperado(d, f):
    """Total geral calculado direto do JSON, sem passar pelo xlsx."""
    qtd = {str(i["id"]): i["quantidade"] for i in d["itens"]}
    total = 0.0
    for k, p in f["precos"].items():
        unit = p.get("unit")
        total += unit * qtd[k] if unit is not None else p["total"]
    imp = 0 if f.get("obs_imposto") == "Incluso" else (f.get("valor_imposto") or 0)  # regra à parte: o exemplo só diz "Incluso" ou "à parte"
    desc = f["valor_desconto"] if f.get("valor_desconto") is not None else total * (f.get("pct_desconto") or 0) / 100
    return total + imp + (f.get("valor_frete") or 0) - desc


class TestMapa(unittest.TestCase):
    def setUp(self):
        with open(ENTRADA, encoding="utf-8") as f:
            self.d = json.load(f)

    def test_exemplo_gera_planilha_que_abre(self):
        ws = gera(self.d)
        self.assertEqual(ws.title, "Cotação")
        self.assertEqual([ws["E4"].value, ws["G4"].value, ws["I4"].value],
                         ["Fornecedor Alfa Ltda", "Fornecedor Beta ME", "Fornecedor Gama Ltda"])
        self.assertEqual([ws.cell(r, 2).value for r in range(8, 12)], [i["descricao"] for i in self.d["itens"]])

    def test_total_geral_confere_com_conta_independente(self):
        ws = gera(self.d)
        n = len(self.d["itens"])
        l_geral = 8 + n + 4  # total, imposto, frete, desconto, depois o total geral
        for i, f in enumerate(self.d["fornecedores"]):
            ref = f"{get_column_letter(5 + 2 * i)}{l_geral}"
            self.assertAlmostEqual(avalia(ws, ref), esperado(self.d, f), places=2, msg=f["nome"])

    def test_item_sem_preco_nao_se_esconde_no_total(self):
        ws = gera(self.d)
        self.assertEqual(ws["I10"].value, "não cotou")
        self.assertEqual(avalia(ws, "I17"), 1)  # linha "Itens sem cotação" do terceiro fornecedor
        self.assertIn("Não cotou: Eletroduto", ws["I18"].value)
        self.assertIn("Cola para eletroduto", ws["I18"].value)

    def test_payload_do_fluxo_n8n_gera_as_mesmas_celulas(self):
        v1 = {
            "codigo_cotacao": self.d["codigo"], "obra": self.d["obra"], "endereco": self.d["endereco"],
            "itens": [{"numero": i["seq"], "id_item_pedido": i["id"], "descricao": i["descricao"],
                       "quantidade": i["quantidade"], "unidade": i["unidade"]} for i in self.d["itens"]],
            "fornecedores": [{**{k: v for k, v in f.items() if k not in ("precos", "contato", "prazo_entrega")},
                              "nome_contato": f["contato"], "prazo_material": f["prazo_entrega"],
                              "itens": {k: {"preco_unit": p.get("unit"), "preco_total": p.get("total")}
                                        for k, p in f["precos"].items()}} for f in self.d["fornecedores"]],
        }
        a, b = gera(self.d), gera(v1)
        celulas = lambda ws: {c.coordinate: c.value for linha in ws.iter_rows() for c in linha if c.value is not None}
        self.assertEqual(celulas(a), celulas(b))

    def test_mais_itens_e_fornecedores_que_o_exemplo(self):
        d = {"obra": "X", "itens": [{"id": k, "seq": k, "descricao": f"item {k}", "quantidade": 2, "unidade": "un"} for k in range(1, 8)],
             "fornecedores": [{"nome": f"F{j}", "precos": {str(k): {"unit": 10 + j} for k in range(1, 8)}} for j in range(5)]}
        ws = gera(d)
        self.assertEqual(ws.max_column, 4 + 2 * 5)
        self.assertAlmostEqual(avalia(ws, f"E{8 + 7 + 4}"), 7 * 2 * 10, places=2)

    def test_sem_itens_e_erro_claro(self):
        with self.assertRaises(ValueError):
            gera({"itens": [], "fornecedores": []})

    def test_menor_total_geral_so_entre_quem_cotou_tudo(self):
        ws = gera(self.d)
        # Alfa (E16) cotou tudo e tem o menor total completo; Gama (I16) tem o menor número, mas deixou um item sem preço
        self.assertTrue(avalia(ws, "I16") < avalia(ws, "E16"))
        self.assertTrue(ws["E16"].fill.start_color.rgb.endswith("C6EFCE"))
        self.assertIsNone(ws["G16"].fill.fill_type)
        self.assertIsNone(ws["I16"].fill.fill_type)

    def test_menor_preco_por_item_e_regra_viva_de_minimo(self):
        ws = gera(self.d)
        regras = {str(cf.sqref): cf.rules[0].formula[0] for cf in ws.conditional_formatting}
        self.assertEqual(sorted(regras), ["E8:E11", "G8:G11", "I8:I11"])
        for faixa_, formula in regras.items():
            self.assertIn("=MIN($E8,$G8,$I8)", formula, faixa_)
            self.assertIn("ISNUMBER", formula, faixa_)

    def test_so_veio_o_total_do_item_unitario_sai_da_conta(self):
        ws = gera(self.d)
        self.assertEqual(ws["G10"].value, 95)  # Beta: total 570 / 6 rolos (exemplo da conversa)
        self.assertEqual(avalia(ws, "H10"), 570)

    def test_imposto_incluso_nao_soma_de_novo_e_silencio_nao_vira_incluso(self):
        def mapa(**extra):
            return gera({"obra": "X", "itens": [{"id": 1, "seq": 1, "descricao": "a", "quantidade": 2, "unidade": "un"}],
                         "fornecedores": [{"nome": "F", "precos": {"1": {"unit": 50}}, **extra}]})
        incluso = mapa(valor_imposto=100, obs_imposto="Incluso")
        self.assertEqual(avalia(incluso, "E13"), 100)  # 2 x 50, o imposto já estava no preço
        a_parte = mapa(valor_imposto=100, obs_imposto="À parte")
        self.assertEqual(avalia(a_parte, "E13"), 200)
        calado = mapa()
        self.assertEqual(calado["F10"].value, "Não informado")  # o fornecedor não disse nada: o mapa não afirma "Incluso"
        self.assertEqual(avalia(calado, "E13"), 100)

    def test_frete_e_desconto(self):
        def mapa(**extra):
            return gera({"obra": "X", "itens": [{"id": 1, "seq": 1, "descricao": "a", "quantidade": 2, "unidade": "un"}],
                         "fornecedores": [{"nome": "F", "precos": {"1": {"unit": 50}}, **extra}]})
        calado = mapa()
        self.assertEqual(calado["F11"].value, "Não informado")  # nada sobre frete: o mapa não inventa modalidade
        self.assertEqual(avalia(calado, "E13"), 100)
        so_valor = mapa(valor_frete=30)
        self.assertIsNone(so_valor["F11"].value)  # tem valor, falta a modalidade
        self.assertEqual(avalia(so_valor, "E13"), 130)
        self.assertEqual(mapa(valor_frete=30, frete="CIF")["F11"].value, "CIF")
        em_percentual = mapa(pct_desconto=5)
        self.assertEqual(avalia(em_percentual, "F12"), 5)  # 5% de 100
        self.assertEqual(avalia(em_percentual, "E12"), 0.05)
        self.assertEqual(avalia(em_percentual, "E13"), 95)
        em_valor = mapa(valor_desconto=7, pct_desconto=5)  # o valor informado vale mais que o percentual
        self.assertEqual(avalia(em_valor, "F12"), 7)
        self.assertEqual(avalia(em_valor, "E13"), 93)
        self.assertEqual(avalia(calado, "E12"), 0)  # sem desconto: 0%

    def test_imposto_incluso_com_negacao_e_acento(self):
        # tabela fixa de frases -> o imposto de R$ 100 entra de novo (200) ou já estava nos preços (100)
        casos = [
            ("Incluso", 100), ("incluso no preço", 100), ("Impostos inclusos", 100), ("ICMS incluso", 100),
            ("incluído no preço", 100), ("INCLUÍDO", 100), ("Incluso, sem destaque", 100), ("sem frete, imposto incluso", 100),
            ("Não incluso", 200), ("nao incluso", 200), ("Não está incluso", 200), ("NÃO ESTÁ INCLUÍDO", 200),
            ("ICMS não está incluso no preço", 200), ("Impostos não inclusos", 200), ("Sem imposto", 200),
            ("À parte", 200), ("Imposto destacado à parte", 200), ("Inclusive frete", 200),
        ]
        for texto, geral in casos:
            with self.subTest(texto):
                ws = gera({"obra": "X", "itens": [{"id": 1, "seq": 1, "descricao": "a", "quantidade": 2, "unidade": "un"}],
                           "fornecedores": [{"nome": "F", "precos": {"1": {"unit": 50}}, "valor_imposto": 100, "obs_imposto": texto}]})
                self.assertEqual(avalia(ws, "E13"), geral)

    def test_destaque_do_menor_total_geral_pede_dois_fornecedores_completos(self):
        itens = [{"id": 1, "seq": 1, "descricao": "a", "quantidade": 1}, {"id": 2, "seq": 2, "descricao": "b", "quantidade": 1}]
        ws = gera({"obra": "X", "itens": itens, "fornecedores": [
            {"nome": "Completo", "precos": {"1": {"unit": 10}, "2": {"unit": 10}}},
            {"nome": "Incompleto", "precos": {"1": {"unit": 1}}}]})
        # com 2 itens o Total Geral fica na linha 14; só um fornecedor cotou tudo, então não há disputa a destacar
        self.assertIsNone(ws["E14"].fill.fill_type)
        self.assertIsNone(ws["G14"].fill.fill_type)

    def test_servidor_http_host_e_port_vazios_valem_como_padrao(self):
        self.assertEqual(servidor_http.endereco({}), ("127.0.0.1", 5000))
        self.assertEqual(servidor_http.endereco({"HOST": "", "PORT": ""}), ("127.0.0.1", 5000))  # .env copiado do .env.example
        self.assertEqual(servidor_http.endereco({"HOST": "0.0.0.0", "PORT": "5077"}), ("0.0.0.0", 5077))

    def test_total_informado_que_difere_de_unit_vezes_qtd_avisa_na_observacao(self):
        ws = gera({"obra": "X", "itens": [{"id": 1, "seq": 1, "descricao": "a", "quantidade": 2, "unidade": "un"}],
                   "fornecedores": [{"nome": "F", "precos": {"1": {"unit": 50, "total": 90}}}]})
        self.assertIn("difere", ws["E15"].value)  # com 1 item: total 9, imposto 10, frete 11, desconto 12, geral 13, sem cotação 14, observação 15
        self.assertIn("90.00", ws["E15"].value)
        self.assertEqual(avalia(ws, "E13"), 100)  # vale unit x qtd; o total divergente só avisa

    def test_itens_fora_de_ordem_saem_ordenados_por_seq(self):
        itens = [{"id": 9, "seq": 3, "descricao": "c", "quantidade": 1, "unidade": "un"},
                 {"id": 7, "seq": 1, "descricao": "a", "quantidade": 1, "unidade": "un"},
                 {"id": 8, "seq": 2, "descricao": "b", "quantidade": 1, "unidade": "un"}]
        ws = gera({"obra": "X", "itens": itens,
                   "fornecedores": [{"nome": "F", "precos": {"7": {"unit": 10}, "8": {"unit": 20}, "9": {"unit": 30}}}]})
        self.assertEqual([ws.cell(r, 2).value for r in (8, 9, 10)], ["a", "b", "c"])
        self.assertEqual([ws.cell(r, 5).value for r in (8, 9, 10)], [10, 20, 30])  # o preço acompanha o item

    def test_seq_misturando_texto_e_numero_nao_quebra(self):
        ws = gera({"obra": "X", "itens": [{"id": 1, "seq": "2", "descricao": "b", "quantidade": 1},
                                           {"id": 2, "seq": 1, "descricao": "a", "quantidade": 1}],
                   "fornecedores": [{"nome": "F", "precos": {"1": {"unit": 1}, "2": {"unit": 1}}}]})
        self.assertEqual([ws.cell(r, 1).value for r in (8, 9)], [1, 2])
        self.assertEqual([ws.cell(r, 2).value for r in (8, 9)], ["a", "b"])

    def test_nao_cotou_fica_marcado_em_vermelho(self):
        ws = gera(self.d)
        self.assertEqual(ws["I10"].value, "não cotou")
        self.assertTrue(ws["I10"].fill.start_color.rgb.endswith("F8CBAD"))
        self.assertIsNone(ws["E10"].fill.fill_type)  # quem cotou não é marcado

    def test_resumo_de_negociacao_traz_as_condicoes_de_cada_fornecedor(self):
        ws = gera(self.d)
        n = len(self.d["itens"])
        pgto, execucao, entrega, validade = (n + 17 + k for k in range(4))
        self.assertEqual([ws.cell(r, 1).value for r in (pgto, execucao, entrega, validade)],
                         ["Condições de pagamento", "Prazo para execução do serviço",
                          "Prazo para entrega do material", "Validade da proposta"])
        for i, f in enumerate(self.d["fornecedores"]):
            c = get_column_letter(5 + 2 * i)
            self.assertEqual(ws[f"{c}{pgto}"].value, f["forma_pagamento"], f["nome"])
            self.assertEqual(ws[f"{c}{execucao}"].value, "Não se aplica", f["nome"])
            self.assertEqual(ws[f"{c}{entrega}"].value, f["prazo_entrega"], f["nome"])
            self.assertEqual(ws[f"{c}{validade}"].value, f["validade"], f["nome"])

    def test_preco_com_virgula_brasileira_e_lido_e_texto_que_nao_e_numero_da_erro(self):
        def mapa(unit):
            return gera({"obra": "X", "itens": [{"id": 1, "seq": 1, "descricao": "a", "quantidade": 2, "unidade": "un"}],
                         "fornecedores": [{"nome": "F", "precos": {"1": {"unit": unit}}}]})
        self.assertEqual(mapa("5,5")["E8"].value, 5.5)  # antes virava "não cotou" sem aviso
        self.assertEqual(avalia(mapa("1.234,50"), "F8"), 2469.0)
        self.assertEqual(mapa("5.5")["E8"].value, 5.5)  # ponto decimal sem vírgula continua valendo
        with self.assertRaises(ValueError):
            mapa("cinco")

    def sobe_servidor(self):
        servidor = servidor_http.ThreadingHTTPServer(("127.0.0.1", 0), servidor_http.Handler)
        threading.Thread(target=servidor.serve_forever, daemon=True).start()
        self.addCleanup(servidor.server_close)
        self.addCleanup(servidor.shutdown)
        return servidor

    def test_servidor_http_responde_400_a_json_que_nao_e_objeto(self):
        servidor = self.sobe_servidor()

        def posta(corpo):
            req = urllib.request.Request(f"http://127.0.0.1:{servidor.server_port}/gerar-mapa", data=corpo, method="POST")
            try:
                with urllib.request.urlopen(req, timeout=10) as r:
                    return r.status, r.read()
            except urllib.error.HTTPError as e:
                return e.code, e.read()

        for ruim in (b"[]", b"[1]", b'"texto"', b"nao e json", b'{"itens": [1]}', b'{"itens": [{"seq": "x"}]}'):
            self.assertEqual(posta(ruim)[0], 400, ruim)
        with open(ENTRADA, "rb") as f:
            status, xlsx = posta(f.read())
        self.assertEqual(status, 200)
        self.assertEqual(xlsx[:2], b"PK")  # xlsx é um zip

    def test_servidor_http_health_404_e_limite_de_tamanho(self):
        porta = self.sobe_servidor().server_port

        def pede(metodo, caminho, content_length=None):
            conn = http.client.HTTPConnection("127.0.0.1", porta, timeout=5)
            self.addCleanup(conn.close)
            conn.putrequest(metodo, caminho)
            if content_length is not None:  # só o cabeçalho: o servidor tem que recusar sem esperar o corpo
                conn.putheader("Content-Length", str(content_length))
            conn.endheaders()
            r = conn.getresponse()
            return r.status, r.read()

        self.assertEqual(pede("GET", "/health"), (200, b'{"ok":true}'))
        self.assertEqual(pede("GET", "/nao-existe")[0], 404)
        self.assertEqual(pede("POST", "/nao-existe", 0)[0], 404)
        self.assertEqual(pede("POST", "/gerar-mapa", servidor_http.LIMITE + 1)[0], 413)  # maior que 5 MB
        self.assertEqual(pede("POST", "/gerar-mapa", 0)[0], 413)  # corpo vazio

if __name__ == "__main__":
    unittest.main()
